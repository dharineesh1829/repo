"""
Helper script to generate .docx versions of sample files for testing file uploads.
"""
from docx import Document
from pathlib import Path

base_dir = Path(__file__).parent

# Resume docx
resume_txt = (base_dir / "sample_resume.txt").read_text(encoding="utf-8")
doc_resume = Document()
doc_resume.add_heading("Alex Chen - Resume", level=0)
for line in resume_txt.splitlines():
    if line.isupper() and len(line) < 40 and not "@" in line:
        doc_resume.add_heading(line, level=1)
    elif line.strip():
        doc_resume.add_paragraph(line)
doc_resume.save(base_dir / "sample_resume.docx")

# JD docx
jd_txt = (base_dir / "sample_job_description.txt").read_text(encoding="utf-8")
doc_jd = Document()
doc_jd.add_heading("Machine Learning & Software Engineer - Job Description", level=0)
for line in jd_txt.splitlines():
    if line.isupper() and len(line) < 40:
        doc_jd.add_heading(line, level=1)
    elif line.strip():
        doc_jd.add_paragraph(line)
doc_jd.save(base_dir / "sample_job_description.docx")

print("Generated sample_resume.docx and sample_job_description.docx successfully.")
